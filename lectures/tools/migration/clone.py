"""Copy slides between .pptx files with python-pptx, the way PowerPoint's
"keep source formatting" paste does: every part a slide depends on (pictures,
charts, media, its layout, master and theme) is copied, rIds are kept, and part
names are renumbered so nothing collides."""
import copy, re
from lxml import etree
from pptx.opc.constants import RELATIONSHIP_TYPE as RT, RELATIONSHIP_TARGET_MODE as RTM
from pptx.opc.package import PartFactory, _Relationship
from pptx.opc.packuri import PackURI

P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"
SKIP = {RT.NOTES_SLIDE}


class Cloner:
    def __init__(self, dst):
        self.dst = dst
        self.pkg = dst.part.package
        self.memo = {}           # id(src part) -> dst part
        self.taken = {str(p.partname) for p in self.pkg.iter_parts()}
        self.layout_for = {}     # id(src layout part) -> dst layout part (same-deck mapping)

    def name(self, partname):
        m = re.match(r"^(.*?)(\d*)(\.\w+)$", str(partname))
        stem, ext = m.group(1), m.group(3)
        self.taken |= {str(p.partname) for p in self.pkg.iter_parts()}
        n = 1
        while f"{stem}{n}{ext}" in self.taken:
            n += 1
        new = f"{stem}{n}{ext}"
        self.taken.add(new)
        return PackURI(new)

    def part(self, src):
        if id(src) in self.memo:
            return self.memo[id(src)]
        new = PartFactory(self.name(src.partname), src.content_type, self.pkg, src.blob)
        self.memo[id(src)] = new
        self.copy_rels(src, new)
        return new

    def copy_rels(self, src, new, skip=()):
        for rId, rel in src.rels.items():
            if rel.reltype in skip:
                continue
            if rel.is_external:
                target, mode = rel.target_ref, RTM.EXTERNAL
            else:
                target, mode = self.target(rel), RTM.INTERNAL
            new.rels._rels[rId] = _Relationship(new.partname.baseURI, rId, rel.reltype, mode, target)

    def target(self, rel):
        t = rel.target_part
        if rel.reltype == RT.SLIDE_LAYOUT and id(t) in self.layout_for:
            return self.layout_for[id(t)]
        if rel.reltype == RT.SLIDE_MASTER and id(t) in self.memo:
            return self.memo[id(t)]
        return self.part(t)

    def register_master(self, master):
        prs_part = self.dst.part
        lst = prs_part._element.find(f"{P}sldMasterIdLst")
        if any(prs_part.related_part(e.get(f"{R}id")) is master for e in lst):
            return
        ids = [int(e.get("id")) for e in lst]
        for m in prs_part.package.iter_parts():
            el = getattr(m, "_element", None)
            if el is not None and el.tag == f"{P}sldMaster":
                ids += [int(e.get("id")) for e in el.iter(f"{P}sldLayoutId")]
        nxt = max(ids + [2147483648]) + 1
        for e in master._element.iter(f"{P}sldLayoutId"):
            e.set("id", str(nxt)); nxt += 1
        rId = prs_part.relate_to(master, RT.SLIDE_MASTER)
        el = etree.SubElement(lst, f"{P}sldMasterId")
        el.set("id", str(nxt)); el.set(f"{R}id", rId)

    def same_deck(self, src_prs):
        """Slides from the deck dst was opened from map onto dst's own layouts by partname."""
        dst_layouts = {str(l.part.partname): l.part for m in self.dst.slide_masters for l in m.slide_layouts}
        for m in src_prs.slide_masters:
            for l in m.slide_layouts:
                if str(l.part.partname) in dst_layouts:
                    self.layout_for[id(l.part)] = dst_layouts[str(l.part.partname)]

    def slide(self, src_slide):
        """Append a copy of src_slide (from any deck) to dst; returns the new slide."""
        src = src_slide.part
        new = PartFactory(self.name(PackURI("/ppt/slides/slide1.xml")), src.content_type, self.pkg, src.blob)
        self.copy_rels(src, new, skip=SKIP)
        layout = new.part_related_by(RT.SLIDE_LAYOUT)
        self.register_master(layout.part_related_by(RT.SLIDE_MASTER))   # no-op for dst's own
        prs_part = self.dst.part
        rId = prs_part.relate_to(new, RT.SLIDE)
        lst = self.dst.slides._sldIdLst
        ids = [int(s.get("id")) for s in lst] or [255]
        el = etree.SubElement(lst, f"{P}sldId")
        el.set("id", str(max(ids) + 1)); el.set(f"{R}id", rId)
        slide = new.slide
        if src_slide.has_notes_slide:
            sn = src_slide.notes_slide.notes_text_frame
            if sn is not None and sn.text.strip():
                dn = slide.notes_slide.notes_text_frame
                dn._txBody.getparent().replace(dn._txBody, copy.deepcopy(sn._txBody))
        if src_slide._element.get("show") == "0":
            slide._element.set("show", "0")
        return slide


def clear_slides(prs):
    lst = prs.slides._sldIdLst
    for s in list(lst):
        prs.part.drop_rel(s.get(f"{R}id"))
        lst.remove(s)
