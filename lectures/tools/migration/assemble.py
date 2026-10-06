"""One-off migration (6 Oct 2026): build one master .pptx per lecture from the qmd files.

Kept for the record; the masters it made are now in lectures/pptx/ and the qmd files are
generated from them by tools/sync.py. Sources it read are now in Drive under 2026_Lectures/_archive/.


Each slide of the qmd either names its source (<!-- src: deck.pptx slide N -->), in which
case the original PowerPoint slide is copied and its text patched to match the qmd, or was
written in Markdown, in which case a native slide is built from the Markdown.
"""
import re, sys, json, shutil, tempfile
from pathlib import Path
from pptx import Presentation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
sys.path.insert(0, str(HERE.parent))
import pptx2qmd as conv
from clone import Cloner, clear_slides

LEC = HERE / "lectures"   # copy of the old lectures/ dir (qmds from Drive _archive + their img/) to run it again
DRIVE = Path.home() / "Google Drive/Work/UCL/Teaching/BASC0005 QM2/2026_Lectures"
SOURCES = {p.name: p for d in ("_archive/Source decks (pre-2026)", "_archive/Drafts (Sep 2026)") for p in (DRIVE / d).glob("*.pptx")}

QMDS = {"w01": LEC / "private/w01.qmd", "w09_private": LEC / "private/w09_private.qmd",
        **{f"w{n:02d}": LEC / f"w{n:02d}.qmd" for n in range(2, 11)}}


def split_slides(text):
    """[(heading, body)] in order; body includes everything up to the next heading."""
    fm, body = re.match(r"^---\n(.*?)\n---\n(.*)$", text, re.S).groups()
    title = re.search(r'^title: "(.*)"', fm, re.M).group(1)
    out, cur, fence = [], None, False
    for line in body.split("\n"):
        if line.startswith("```"):
            fence = not fence
        if not fence and re.match(r"^##? ", line) or (not fence and line == "##"):
            cur = [line, []]; out.append(cur)
        elif cur is not None:
            cur[1].append(line)
    return title, [(h, "\n".join(b).strip("\n")) for h, b in out]


def src_of(body):
    m = re.search(r"<!-- src: (.*?\.pptx) slide (\d+) -->", body)
    return (m.group(1), int(m.group(2))) if m else None


_PRS = {}
def deck(name):
    if name not in _PRS:
        _PRS[name] = Presentation(SOURCES[name])
    return _PRS[name]


def norm(chunk):
    """Comparable form of a slide's Markdown."""
    t = re.sub(r"<!--.*?-->", "", chunk, flags=re.S)
    t = re.sub(r"\(img/[^)]+\)", "(IMG)", t)
    t = re.sub(r'background-image="[^"]+"', 'background-image="IMG"', t)
    t = re.sub(r" ?\{\.(center|smaller)\}", "", t)
    t = re.sub(r"^(#+ .*)$", lambda m: re.sub(r"\s+", " ", m.group(1).replace("\x0b", " ")), t, flags=re.M)
    t = re.sub(r"\n{2,}", "\n\n", t)
    return "\n".join(l.rstrip() for l in t.strip().split("\n"))


def convert_one(slide, i, tmp):
    conv.SW, conv.SH = 12192000, 6858000
    ctx = dict(img_dir=tmp, rel="img/x", fallback=lambda i: "img/x/FALLBACK.jpg")
    md, kind, why = conv.slide_md(slide, i, ctx)
    return md, kind



import difflib, copy, copy as _copy
from pptx.oxml.ns import qn
from mdslide import latex, build_slide, set_runs, split_notes, inline, layout as find_layout

MARK = re.compile(r"^\s*(?:- |\d+\. )?")


def strip_mark(l):
    return MARK.sub("", l, count=1)


def paragraphs(slide):
    """(paragraph, kind) for every paragraph a converted slide can show."""
    out = []
    def walk(shapes):
        for sh in shapes:
            if sh.shape_type == 6:                 # group
                walk(sh.shapes); continue
            if getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for c in row.cells:
                        out.extend((p, "cell") for p in c.text_frame.paragraphs)
            elif sh.has_text_frame:
                kind = "title" if conv.is_title(sh) else "text"
                out.extend((p, kind) for p in sh.text_frame.paragraphs)
    walk(slide.shapes)
    return out


def body_lines(md):
    body, notes = split_notes_any(md)
    lines = [l for l in body.split("\n") if l.strip() and not l.startswith(("![", ":::", "<!--"))]
    return lines, notes


def split_notes_any(md):
    m = re.search(r"::: notes\s*(.*?)\n:::\s*$", md, re.S)
    return (md[:m.start()], m.group(1).strip()) if m else (md, "")


def find_para(paras, want, used, plain=False, raw=None):
    for w in ([want] + ([raw.strip()] if raw else [])):
        hit = _find(paras, w, used, plain)
        if hit[0] is not None:
            return hit
    return None, None


def _find(paras, want, used, plain):
    for k, (p, kind) in enumerate(paras):
        if k in used:
            continue
        have = re.sub(r"\s+", " ", p.text if plain else conv.md_runs(p)).strip()
        w = re.sub(r"\s+", " ", want).strip()
        if have and (have == w or strip_mark(have) == w):
            used.add(k); return p, kind
    return None, None


def drop_row(slide, want):
    for sh in slide.shapes:
        if getattr(sh, "has_table", False) and sh.has_table:
            for tr in sh.table._tbl.tr_lst:
                cells = ["".join(t.text or "" for t in tc.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t")).strip() for tc in tr.tc_lst]
                if [re.sub(r"\*", "", c) for c in want] == cells:
                    tr.getparent().remove(tr); return True
    return False


def code_patch(slide, target_md):
    """Replace a code text box's lines with the qmd's code block."""
    m = re.search(r"```[^\n]*\n(.*?)\n```", target_md, re.S)
    if not m:
        return False
    for sh in slide.shapes:
        if sh.has_text_frame and conv.code_block(sh.text_frame):
            tf = sh.text_frame
            rprs = [r._r.find(qn("a:rPr")) for p in tf.paragraphs for r in p.runs]
            rpr = next((x for x in rprs if x is not None), None)
            font = next((r.font.name for p in tf.paragraphs for r in p.runs if r.font.name), "Consolas")
            ppr = tf.paragraphs[0]._p.find(qn("a:pPr"))
            for p in list(tf.paragraphs)[1:]:
                p._p.getparent().remove(p._p)
            for k, line in enumerate(m.group(1).split("\n")):
                p = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                if k and ppr is not None: p._p.insert(0, copy.deepcopy(ppr))
                for r in list(p._p):
                    if r.tag in (qn("a:r"), qn("a:br")): p._p.remove(r)
                r = p.add_run(); r.text = line
                if rpr is not None:
                    old = r._r.find(qn("a:rPr"))
                    if old is not None: r._r.remove(old)
                    r._r.insert(0, copy.deepcopy(rpr))
                r.font.name = font
            return True
    return False


def patch(slide, conv_md, target_md, log):
    if code_patch(slide, target_md):
        conv_md = convert_one(slide, 0, Path(tempfile.mkdtemp()))[0]
    paras = paragraphs(slide)
    a, an = body_lines(norm(conv_md)); b, bn = body_lines(norm(target_md))
    # lines on both sides (ignoring bold, in any order) need no patch: reordering is layout, not text
    key = lambda l: l.replace("**", "").strip()
    ka, kb = {key(l) for l in a}, {key(l) for l in b}
    b_all = b
    a = [l for l in a if key(l) not in kb]
    b = [l for l in b if key(l) not in ka]
    used = set()
    # table rows: diff cell by cell
    def cells(l): return [c.strip() for c in re.split(r"(?<!\\)\|", l)[1:-1]]
    sm = difflib.SequenceMatcher(None, a, b, autojunk=False)
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == "equal":
            continue
        olds, news = a[i1:i2], b[j1:j2]
        for k in range(max(len(olds), len(news))):
            old = olds[k] if k < len(olds) else None
            new = news[k] if k < len(news) else None
            if old and new and old.startswith("|") and new.startswith("|"):
                for co, cn in zip(cells(old), cells(new)):
                    if co != cn:
                        p, _ = find_para(paras, co, used)
                        if p is None: log.append(f"cell? {co!r}"); continue
                        set_runs(p, cn)
                continue
            if old is not None:
                want = re.sub(r"^#+ ", "", old) if old.startswith("#") else strip_mark(old)
                p, kind = find_para(paras, want, used, raw=re.sub(r"^- ", "", old))
                if p is None and old.startswith("|") and new is None:      # drop a table row
                    if drop_row(slide, cells(old)): continue
                if p is None:
                    log.append(f"line? {old!r} -> {new!r}"); continue
                if new is None:
                    if kind == "title":
                        log.append(f"title delete? {old!r}"); continue
                    p._p.getparent().remove(p._p); continue
                newt = re.sub(r"^#+ ", "", new) if new.startswith("#") else strip_mark(new)
                newt = re.sub(r"\s*\{[^}]*\}$", "", newt) if new.startswith("#") else newt
                set_runs(p, newt)
                last = p
            else:                                     # insertion after the previous line
                nb = b_all.index(new) if new in b_all else 0
                prev = b_all[nb - 1] if nb > 0 else None
                anchor = None
                if prev is not None:
                    want = strip_mark(prev)
                    for q, kind in paras:
                        if re.sub(r"\s+", " ", conv.md_runs(q)).strip() == want.strip():
                            anchor = q
                    if anchor is None:
                        for q, kind in paras:
                            if conv.md_runs(q).strip() == strip_mark(new).strip():
                                anchor = None; break
                if anchor is None:
                    log.append(f"insert? {new!r}"); continue
                el = _copy.deepcopy(anchor._p)
                anchor._p.addnext(el)
                from pptx.text.text import _Paragraph
                np_ = _Paragraph(el, anchor._parent)
                lvl = (len(new) - len(new.lstrip())) // 4
                if np_.level != lvl:
                    np_.level = lvl
                    pPr = el.find(qn("a:pPr"))
                    for k in ("marL", "indent"):
                        if pPr is not None: pPr.attrib.pop(k, None)
                set_runs(np_, strip_mark(new))
                paras.append((np_, "text"))
    # notes and, for slides kept as pictures, the slide text
    st_a = re.search(r"Slide text: (.*)$", an, re.S); st_b = re.search(r"Slide text: (.*)$", bn, re.S)
    if st_a and st_b and st_a.group(1).strip() != st_b.group(1).strip():
        sa = [x.strip() for x in st_a.group(1).split(" · ")]; sb = [x.strip() for x in st_b.group(1).split(" · ")]
        sm = difflib.SequenceMatcher(None, sa, sb, autojunk=False)
        for op, i1, i2, j1, j2 in sm.get_opcodes():
            if op == "replace" and i2 - i1 == j2 - j1:
                for o, n in zip(sa[i1:i2], sb[j1:j2]):
                    p, _ = find_para(paras, o, used, plain=True)
                    if p is None: log.append(f"slide text? {o!r}"); continue
                    set_runs(p, n.replace("*", r"\*"))
            elif op == "delete":
                for o in sa[i1:i2]:
                    p, _ = find_para(paras, o, used, plain=True)
                    if p is None: log.append(f"slide text del? {o!r}"); continue
                    p._p.getparent().remove(p._p)
            elif op != "equal":
                log.append(f"slide text {op} {sa[i1:i2]} -> {sb[j1:j2]}")
    na = re.sub(r"\n*Slide text: .*$", "", an, flags=re.S).strip()
    nb = re.sub(r"\n*Slide text: .*$", "", bn, flags=re.S).strip()
    if na != nb:
        slide.notes_slide.notes_text_frame.text = nb


KEEP = re.compile(r"^(‹#›|Ollie Ballinger.*|Quantitative Methods II.*|Click.*|Edit Master.*|.*Second level.*|)$", re.S)


def fix_masters(prs, topic):
    """Masters carry hard-coded footers from their old decks: the old topic, '/41'-style
    slide totals and dates. Point the topic at this lecture and drop the stale rest."""
    for m in prs.slide_masters:
        for part in [m] + list(m.slide_layouts):
            for sh in part.shapes:
                if not sh.has_text_frame:
                    continue
                full = sh.text_frame.text.strip()
                if sh.is_placeholder and sh.placeholder_format.type in (1, 2, 3, 4, 7):
                    continue                          # title/body placeholders: prompt text
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if re.fullmatch(r"\s*/\s*\d+\s*", r.text) or re.fullmatch(r"\d{1,2}/\d{1,2}/\d{2,4}", r.text.strip()):
                            r.text = ""
                if full and not KEEP.match(full) and not re.fullmatch(r"[\d/\s]+|.*#.*", full):
                    runs = [r for p in sh.text_frame.paragraphs for r in p.runs]
                    if runs:
                        runs[0].text = topic
                        for r in runs[1:]: r.text = ""


def w01_geofence(prs):
    """The geofencing slide: the commented code no longer fits beside the data table."""
    from pptx.util import Inches, Pt
    s = next(x for x in prs.slides if x.shapes.title is not None and x.shapes.title.text.startswith("Geofencing"))
    code = next(sh for sh in s.shapes if sh.has_text_frame and conv.code_block(sh.text_frame))
    table = next(sh for sh in s.shapes if getattr(sh, "has_table", False) and sh.has_table)
    code.left, code.top, code.width, code.height = Inches(0.9), Inches(1.25), Inches(11.5), Inches(2.6)
    for p in code.text_frame.paragraphs:
        for r in p.runs: r.font.size = Pt(18)
    lab = s.shapes.add_textbox(Inches(0.9), Inches(4.0), Inches(3), Inches(0.4))
    r = lab.text_frame.paragraphs[0].add_run(); r.text = "data ="; r.font.name = "Consolas"; r.font.size = Pt(18)
    table.left, table.top, table.width = Inches(0.9), Inches(4.5), Inches(11.5)
    for row in table.table.rows:
        for c in row.cells:
            for p in c.text_frame.paragraphs:
                for r in p.runs: r.font.size = Pt(11)


TWEAKS = {"w01": [w01_geofence]}


def base_deck(slides):
    counts = {}
    for h, b in slides:
        s = src_of(b)
        if s: counts[s[0]] = counts.get(s[0], 0) + 1
    for name in sorted(counts, key=counts.get, reverse=True):
        names = [l.name for m in deck(name).slide_masters for l in m.slide_layouts]
        if any(n in names for n in ("Section Header", "SECTION_HEADER")):
            return name
    raise SystemExit("no base deck")


def build(wk, out, tmp):
    qmd = QMDS[wk]
    title, slides = split_slides(qmd.read_text())
    base = base_deck(slides)
    prs = Presentation(SOURCES[base])
    clear_slides(prs)
    cl = Cloner(prs); cl.same_deck(deck(base))
    t = prs.slides.add_slide(find_layout(prs, "Title Slide"))
    t.shapes.title.text = title
    for ph in t.placeholders:
        if ph.placeholder_format.idx == 1:
            ph.text = "BASC0005 Quantitative Methods 2\nOllie Ballinger"
    log = {}
    for k, (h, b) in enumerate(slides, 2):
        s = src_of(b)
        target = h + "\n" + b
        if s:
            new = cl.slide(deck(s[0]).slides[s[1] - 1])
            md, _ = convert_one(new, k, tmp)
            if norm(md) != norm(target):
                L = []
                patch(new, md, target, L)
                if L: log[k] = L
        else:
            body, notes = split_notes(b)
            build_slide(prs, h, body.strip(), qmd.parent, notes)
    for fn in TWEAKS.get(wk, []):
        fn(prs)
    fix_masters(prs, title.split(":")[0].strip())
    prs.save(out)
    return title, slides, log, base


def verify(path, slides, tmp):
    prs = Presentation(path)
    bad = []
    for k, ((h, b), s) in enumerate(zip(slides, list(prs.slides)[1:]), 2):
        md, kind = convert_one(s, k, tmp)
        a, an = body_lines(norm(md)); t, tn = body_lines(norm(h + "\n" + b))
        flat = lambda L: [re.sub(r"\*|\s+", "", strip_mark(re.sub(r"\$([^$]+)\$", lambda m: latex(m.group(1)), x.replace("$$", "$")))) for x in L]
        flat2 = lambda L: [x for x in flat(L) if not re.fullmatch(r"\|[-|:]*\|?", x)]
        if flat2(a) != flat2(t):
            bad.append((k, kind, a, t))
    return bad, len(prs.slides) - 1


if __name__ == "__main__":
    tmp = Path(tempfile.mkdtemp())
    outdir = HERE / "masters"; outdir.mkdir(exist_ok=True)
    for wk in [w for w in sys.argv[1:] if not w.startswith("-")] or QMDS:
        out = outdir / f"{wk}.pptx"
        title, slides, log, base = build(wk, out, tmp)
        bad, n = verify(out, slides, tmp)
        print(f"{wk}: base={base} slides={n}/{len(slides)} patch-issues={len(log)} mismatches={len(bad)}")
        for k, L in log.items():
            for x in L: print(f"   patch {k}: {x[:150]}")
        if "-v" in sys.argv:
            for k, kind, a, t in bad:
                d = [l for l in difflib.unified_diff(a, t, lineterm="", n=0) if not l.startswith(("---", "+++", "@@"))]
                print(f"   slide {k} ({kind}):"); print("\n".join("      " + x[:150] for x in d[:10]))
