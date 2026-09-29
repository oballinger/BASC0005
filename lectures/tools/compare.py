"""Contact sheets for checking a conversion: original slide (left) vs rendered reveal.js slide (right).

usage: compare.py deck.pptx orig.pdf deck.qmd deck.html out_prefix [native|image|all]
"""
import io, re, sys
from pathlib import Path
from PIL import Image, ImageDraw
import pypdfium2 as pdfium
from pptx import Presentation
from playwright.sync_api import sync_playwright

pptx, pdf, qmd, html, outp = sys.argv[1:6]
only = sys.argv[6] if len(sys.argv) > 6 else "native"

# visible slide number -> PDF page (LibreOffice skips hidden slides)
page_of, n = {}, 0
for i, s in enumerate(Presentation(pptx).slides, 1):
    if s._element.get("show") != "0":
        n += 1; page_of[i] = n

src = [(int(m.group(1)), "image" if "fallback" in m.group(2) else "native")
       for m in re.finditer(r"<!-- src: .*? slide (\d+) -->\n(##? [^\n]*)", Path(qmd).read_text())]
doc = pdfium.PdfDocument(pdf)
W, H = 800, 450
shots = []
with sync_playwright() as p:
    b = p.chromium.launch(channel="chrome")
    pg = b.new_page(viewport={"width": 1600, "height": 900})
    pg.goto("file://" + str(Path(html).resolve())); pg.wait_for_timeout(1500)
    for k, (sn, kind) in enumerate(src):
        pg.evaluate("Reveal.next()"); pg.wait_for_timeout(300)  # linear order, skips the title slide
        if only != "all" and kind != only:
            continue
        shots.append((k + 1, sn, kind, Image.open(io.BytesIO(pg.screenshot())).convert("RGB")))
    b.close()

Path(outp).parent.mkdir(parents=True, exist_ok=True)
for g in range(0, len(shots), 6):
    grp = shots[g:g + 6]
    sheet = Image.new("RGB", (W * 2 + 10, (H + 24) * len(grp)), "white")
    d = ImageDraw.Draw(sheet)
    for r, (k, sn, kind, shot) in enumerate(grp):
        y = r * (H + 24)
        d.text((5, y + 5), f"new slide {k}  <-  source slide {sn} ({kind})", fill="black")
        pgi = doc[page_of[sn] - 1]
        orig = pgi.render(scale=W / pgi.get_width()).to_pil().convert("RGB").resize((W, H))
        sheet.paste(orig, (0, y + 22)); sheet.paste(shot.resize((W, H)), (W + 10, y + 22))
    sheet.save(f"{outp}_{g // 6:02d}.jpg", quality=80)
print(f"{len(shots)} slides -> {(len(shots) + 5) // 6} sheets")
