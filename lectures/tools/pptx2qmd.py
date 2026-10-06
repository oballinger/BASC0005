"""Convert a .pptx lecture into a Quarto reveal.js .qmd.

Slides built from titles, bullets and pictures become native Markdown.
Slides with things Markdown can't express (diagrams, grouped shapes, charts,
text laid over images) fall back to a rendered image of the slide, with the
slide text kept in the speaker notes so it stays searchable.

usage: python pptx2qmd.py deck.pptx rendered.pdf out_dir slug
(normally run via sync.py, which also exports the PDF and renders the site)
"""
import argparse, hashlib, io, re, subprocess, sys
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
from pptx.util import Emu

MAXW = 1920  # longest image edge on the web
SW = SH = None


def hidden(slide):
    return slide._element.get("show") == "0"


def parse_ranges(spec, n):
    if not spec:
        return list(range(1, n + 1))
    out = []
    for part in spec.split(","):
        a, _, b = part.partition("-")
        out += list(range(int(a), int(b or a) + 1))
    return out


def is_title(sh):
    return sh.is_placeholder and sh.placeholder_format.type in (
        PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)


def is_chrome(sh):
    """Slide numbers, footers and dates: never content."""
    if sh.is_placeholder and sh.placeholder_format.type in (
            PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE):
        return True
    return sh.has_text_frame and sh.text_frame.text.strip() in ("‹#›", "<#>")


def md_runs(par):
    out = []
    for r in par.runs:
        t = r.text.replace("*", r"\*")
        if not t.strip():
            out.append(t); continue
        lead, core, trail = re.match(r"^(\s*)(.*?)(\s*)$", t, re.S).groups()
        if r.font.bold:
            core = f"**{core}**"
        elif r.font.italic:
            core = f"*{core}*"
        if r.hyperlink and r.hyperlink.address:
            core = f"[{core}]({r.hyperlink.address})"
        out.append(lead + core + trail)
    return "".join(out).strip()


def no_bullet(p):
    pPr = p._p.find("{http://schemas.openxmlformats.org/drawingml/2006/main}pPr")
    return pPr is not None and pPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}buNone") is not None


def text_block(tf):
    """Paragraphs as Markdown: bullets and numbers as lists, explicit no-bullet paragraphs as text."""
    lines = []
    for p in tf.paragraphs:
        s = md_runs(p)
        if not s:
            continue
        if no_bullet(p) and p.level == 0:
            if p.alignment == 2:                   # centred (e.g. an equation on its own line)
                s = f'[{s}]{{style="display:block;text-align:center"}}'
            lines.append(("plain", s)); continue
        numbered = p._p.find(".//{http://schemas.openxmlformats.org/drawingml/2006/main}buAutoNum") is not None
        lines.append(("item", "    " * p.level + ("1. " if numbered else "- ") + s))
    # a single short paragraph reads better as plain text than a lone bullet
    if len(lines) == 1:
        return re.sub(r"^\s*(- |1\. )", "", lines[0][1])
    out = []
    for k, (kind, s) in enumerate(lines):
        if out and (kind == "plain" or lines[k - 1][0] == "plain"):
            out.append("")                     # Markdown needs a blank line around plain paragraphs
        out.append(s)
    return "\n".join(out)


def save_picture(sh, img_dir, rel):
    """Write the picture (crop applied, downscaled) and return its relative path."""
    blob, ext = sh.image.blob, sh.image.ext.lower()
    if ext in ("wmf", "emf"):
        return None
    try:
        im = Image.open(io.BytesIO(blob))
    except Exception:
        return None
    anim = getattr(im, "is_animated", False)
    h = hashlib.sha1(blob).hexdigest()[:10]
    if anim:  # keep GIF animation as-is
        name = f"{h}.gif"
        (img_dir / name).write_bytes(blob)
        return f"{rel}/{name}"
    cl, cr, ct, cb = sh.crop_left, sh.crop_right, sh.crop_top, sh.crop_bottom
    if any((cl, cr, ct, cb)):
        w, hgt = im.size
        box = (int(w * max(cl, 0)), int(hgt * max(ct, 0)),
               int(w * (1 - max(cr, 0))), int(hgt * (1 - max(cb, 0))))
        if box[2] > box[0] and box[3] > box[1]:
            im = im.crop(box)
            h += "c" + hashlib.sha1(repr(box).encode()).hexdigest()[:4]  # one file per distinct crop
    if max(im.size) > MAXW:
        im.thumbnail((MAXW, MAXW))
    photo = im.mode in ("RGB", "CMYK", "YCbCr") or ext in ("jpg", "jpeg")
    if photo:
        name = f"{h}.jpg"
        im.convert("RGB").save(img_dir / name, quality=85, optimize=True)
    else:
        name = f"{h}.png"
        im.save(img_dir / name, optimize=True)
    return f"{rel}/{name}"


def overlaps(a, b):
    ax, ay, aw, ah = a.left or 0, a.top or 0, a.width or 0, a.height or 0
    bx, by, bw, bh = b.left or 0, b.top or 0, b.width or 0, b.height or 0
    ix = max(0, min(ax + aw, bx + bw) - max(ax, bx))
    iy = max(0, min(ay + ah, by + bh) - max(ay, by))
    small = min(aw * ah, bw * bh) or 1
    return ix * iy / small > 0.3


def classify(slide):
    """Return (native_ok, reason, parts)."""
    title, texts, pics, tables, reasons = None, [], [], [], []
    for sh in slide.shapes:
        if is_chrome(sh):
            continue
        st = sh.shape_type
        if is_title(sh):
            title = sh; continue
        if st == MSO_SHAPE_TYPE.PICTURE or (sh.is_placeholder and hasattr(sh, "image")):
            try:
                sh.image; pics.append(sh); continue
            except Exception:
                reasons.append("empty picture placeholder"); continue
        if sh.has_table if hasattr(sh, "has_table") else False:
            tables.append(sh); continue
        if st in (MSO_SHAPE_TYPE.GROUP, MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.LINE,
                  MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.MEDIA, MSO_SHAPE_TYPE.DIAGRAM,
                  MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT):
            reasons.append(str(st).split(".")[-1].split(" ")[0]); continue
        if st == MSO_SHAPE_TYPE.AUTO_SHAPE:
            filled = False
            try:
                filled = sh.fill.type is not None and sh.fill.type != 5  # 5 = background
            except Exception:
                pass
            if filled or not (sh.has_text_frame and sh.text_frame.text.strip()):
                reasons.append("drawn shape"); continue
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh); continue
        if getattr(sh, "has_chart", False) or st == MSO_SHAPE_TYPE.PLACEHOLDER:
            continue
    xml = slide._element.xml
    if "oMath" in xml or "AlternateContent" in xml:
        reasons.append("equation")
    if any(sh._element.find(".//{http://schemas.openxmlformats.org/presentationml/2006/main}spPr/"
                            "{http://schemas.openxmlformats.org/drawingml/2006/main}blipFill") is not None
           for sh in slide.shapes if sh.shape_type != MSO_SHAPE_TYPE.PICTURE):
        reasons.append("image-filled shape")
    for t in texts:
        if any(overlaps(t, p) for p in pics):
            reasons.append("text over image")
    if len(texts) > 2:
        reasons.append(f"{len(texts)} text boxes")
    return (not reasons), ", ".join(sorted(set(reasons))), (title, texts, pics, tables)


BODY_H = 700  # px of a 1600x900 slide left below the title


def layout(items):
    """Lay out (kind, shape, markdown) items the way they sit on the slide.

    Items whose horizontal extents overlap form a column; columns sit side by side with widths
    in proportion to the slide, items within a column run top to bottom. Pictures get a share of
    the column height in proportion to their height on the slide, and fill their column's width
    (class .pic) instead of showing at their pixel size. Returns None when two pictures sit side
    by side inside one column (a collage), which Markdown columns can't express.
    """
    box = lambda sh: (sh.left or 0, sh.top or 0, sh.width or 0, sh.height or 0)
    cols = []
    for it in sorted(items, key=lambda it: box(it[1])[0]):
        x, _, w, _ = box(it[1])
        cols.append({"l": x, "r": x + w, "items": [it]})
    merged = True
    while merged:  # merge columns whose extents overlap by over half the narrower one
        merged = False
        for a in range(len(cols)):
            for b in range(a + 1, len(cols)):
                A, B = cols[a], cols[b]
                ov = min(A["r"], B["r"]) - max(A["l"], B["l"])
                if ov > 0.5 * min(A["r"] - A["l"], B["r"] - B["l"]):
                    A["l"], A["r"] = min(A["l"], B["l"]), max(A["r"], B["r"])
                    A["items"] += B["items"]; del cols[b]; merged = True; break
            if merged:
                break
    cols.sort(key=lambda c: c["l"])
    for c in cols:
        c["items"].sort(key=lambda it: box(it[1])[1])
        ps = [box(it[1]) for it in c["items"] if it[0] == "pic"]
        for m in range(len(ps)):
            for n in range(m + 1, len(ps)):
                (_, t1, _, h1), (_, t2, _, h2) = ps[m], ps[n]
                if min(t1 + h1, t2 + h2) - max(t1, t2) > 0.3 * min(h1, h2):
                    return None

    def render(c, single):
        tot = sum(box(it[1])[3] for it in c["items"]) or 1
        npic = sum(it[0] == "pic" for it in c["items"])
        out = []
        for kind, sh, md in c["items"]:
            if kind != "pic":
                out.append(md)
            elif single and npic == 1:
                out.append(f"![]({md}){{.r-stretch fig-align=\"center\"}}")
            else:
                h = min(BODY_H, round(BODY_H * box(sh)[3] / tot))
                out.append(f"![]({md}){{.pic style=\"max-height:{h}px\"}}")
        return "\n\n".join(out)

    if len(cols) == 1:
        return render(cols[0], True)
    span = sum(c["r"] - c["l"] for c in cols) or 1
    widths = [round(100 * (c["r"] - c["l"]) / span) for c in cols]
    widths[-1] = 100 - sum(widths[:-1])
    return "::: {.columns}\n" + "\n".join(
        f":::: {{.column width=\"{w}%\"}}\n{render(c, False)}\n::::" for c, w in zip(cols, widths)) + "\n:::"


def table_md(sh):
    rows = [[" ".join(md_runs(p) for p in c.text_frame.paragraphs).replace("|", r"\|").strip()
             for c in r.cells] for r in sh.table.rows]
    if not rows:
        return ""
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    if len(rows[0]) >= 8:                      # wide tables: small type, one line per row
        return '::: {style="font-size:0.42em; white-space:nowrap"}\n' + "\n".join(out) + "\n:::"
    return "\n".join(out)


def slide_text(slide):
    return "\n".join(sh.text_frame.text.strip() for sh in slide.shapes
                     if sh.has_text_frame and sh.text_frame.text.strip() and not is_chrome(sh))


def notes_of(slide):
    if slide.has_notes_slide:
        t = slide.notes_slide.notes_text_frame.text.strip() if slide.notes_slide.notes_text_frame else ""
        return t
    return ""


MONO = ("consolas", "courier", "courier new", "menlo", "monaco", "source code pro",
        "roboto mono", "fira code", "fira mono", "jetbrains mono", "sf mono", "andale mono")


def code_block(tf):
    """A text box set entirely in a monospace font is code: return a fenced block, else None."""
    runs = [r for p in tf.paragraphs for r in p.runs if r.text.strip()]
    if not runs or not all((r.font.name or "").lower() in MONO for r in runs):
        return None
    lines = ["".join(r.text for r in p.runs).replace("\u2018", "'").replace("\u2019", "'")
             .replace("\u201c", '"').replace("\u201d", '"').rstrip() for p in tf.paragraphs]
    while lines and not lines[-1].strip():
        lines.pop()
    return '```{.python code-line-numbers="false"}\n' + "\n".join(lines) + "\n```"


def frame_md(sh):
    return code_block(sh.text_frame) or text_block(sh.text_frame)


def small_type(texts, tables):
    """True when the slide's text is set small in PowerPoint (median explicit size <= 16pt):
    the web version then gets reveal's .smaller class."""
    sizes = []
    frames = [t.text_frame for t in texts] + [c.text_frame for t in tables for r in t.table.rows for c in r.cells]
    for tf in frames:
        for p in tf.paragraphs:
            for r in p.runs:
                if r.font.size and r.text.strip():
                    sizes += [r.font.size.pt] * len(r.text)
    if not sizes:
        return False
    sizes.sort()
    return sizes[len(sizes) // 2] <= 16


def fallback_notes(s, notes):
    txt = slide_text(s)
    return "\n\n".join(x for x in [notes, "Slide text: " + txt.replace("\n", " · ") if txt else ""] if x)


def slide_md(s, i, ctx):
    """Markdown for slide `s` (number `i` in the deck).

    ctx: img_dir (Path) and rel (str) for pictures; fallback(i) -> image path of the rendered slide.
    Returns (markdown, kind, reason).
    """
    img_dir, rel, fallback = ctx["img_dir"], ctx["rel"], ctx["fallback"]
    tag = f"<!-- slide {i} -->"
    ok, why, (tsh, texts, pics, tables) = classify(s)
    ttl = tsh.text_frame.text.strip().replace("\n", " ").replace("\x0b", " ") if tsh is not None else ""
    ttl = re.sub(r"\s+", " ", ttl)
    lname = s.slide_layout.name.lower()
    if ("section" in lname) and ok and not pics and len(texts) <= 1:
        return f"# {ttl or text_block(texts[0].text_frame)}\n{tag}", "native", ""
    notes = notes_of(s)
    if ok:
        pic_paths = [p for p in (save_picture(p, img_dir, rel) for p in pics) if p]
        if len(pic_paths) < len(pics):
            ok, why = False, "unsupported image format"
    body, attrs = [], []
    if ok:
        flow = sorted([(t, frame_md(t)) for t in texts] + [(t, table_md(t)) for t in tables],
                      key=lambda x: (x[0].top or 0, x[0].left or 0))
        txt = [md for t, md in flow if md and t in texts]
        tbl = [md for t, md in flow if md and t in tables]
        if pic_paths and not txt and not tbl and len(pic_paths) == 1:
            p = pics[0]
            if (p.width or 0) * (p.height or 0) > 0.7 * SW * SH and not ttl:
                head = f'## {{background-image="{pic_paths[0]}" background-size="contain"}}'
                return _chunk(head, tag, [], notes), "native", ""
            body.append(f"![]({pic_paths[0]}){{.r-stretch fig-align=\"center\"}}")
        elif pic_paths:
            items = ([("pic", p, pp) for p, pp in zip(pics, pic_paths)]
                     + [("md", t, frame_md(t)) for t in texts]
                     + [("md", t, table_md(t)) for t in tables])
            md = layout([it for it in items if it[2]])
            if md is None:  # pictures tiled or overlapping: keep the slide as drawn
                ok, why = False, "tiled pictures"
            else:
                body.append(md)
        else:
            items = [("md", t, md) for t, md in flow if md]
            md = layout(items) if len(items) > 1 else None
            body.append(md if md else "\n\n".join(m for _, _, m in items))
    if not ok:
        head = f'## {{background-image="{fallback(i)}" background-size="contain" .fallback}}'
        return _chunk(head, tag, [], fallback_notes(s, notes)), "image", why
    body = ["\n\n".join(body)] if body else []
    joined = body[0] if body else ""
    if not joined.strip() and ttl:
        attrs.append(".center")                      # question / title-only slides
    elif small_type(texts, tables):
        attrs.append(".smaller")
    head = (f"## {ttl}" if ttl else "##") + (" {" + " ".join(attrs) + "}" if attrs else "")
    if not ttl and not attrs:
        head = "## {.notitle}"
    return _chunk(head, tag, body, notes), "native", ""


def _chunk(head, tag, body, notes):
    # heading first: an HTML comment before the first heading makes an empty slide
    out = [head, tag, ""] + body
    if notes:
        out += ["", "::: notes", notes, ":::"]
    text = "\n".join(out).rstrip()
    return re.sub(r"([^\n])\n(\|[^\n]*\|\n\|[-|: ]+\|\n)", r"\1\n\n\2", text + "\n").rstrip()


def deck_title(prs):
    s = prs.slides[0]
    if "title" in s.slide_layout.name.lower():
        for sh in s.shapes:
            if is_title(sh) and sh.text_frame.text.strip():
                return re.sub(r"\s+", " ", sh.text_frame.text.strip())
    return None


def pdf_renderer(pdf, prs, img_dir, rel):
    """fallback(i) for convert(): render slide i's page of the deck's PDF export."""
    import pypdfium2 as pdfium
    doc = pdfium.PdfDocument(str(pdf))
    page_of, page = {}, 0
    for i, s in enumerate(prs.slides, 1):        # PDF pages follow visible slides only
        if not hidden(s):
            page += 1; page_of[i] = page

    def render(i):
        pg = doc[page_of[i] - 1]
        im = pg.render(scale=1600 / pg.get_width()).to_pil().convert("RGB")
        im = im.crop((0, 0, im.width, int(im.height * 0.945)))  # drop the old footer strip
        name = f"slide_{i:03d}.jpg"
        im.save(img_dir / name, quality=85, optimize=True)
        return f"{rel}/{name}"
    render.close = doc.close
    return render


def convert(deck, out_dir, slug, fallback=None, pdf=None):
    """Write out_dir/slug.qmd and out_dir/img/slug/ from a deck. Returns [(slide, kind, why)]."""
    global SW, SH
    prs = Presentation(deck)
    SW, SH = prs.slide_width, prs.slide_height
    img_dir = Path(out_dir) / "img" / slug
    img_dir.mkdir(parents=True, exist_ok=True)
    rel = f"img/{slug}"
    if fallback is None:
        fallback = pdf_renderer(pdf, prs, img_dir, rel)
    title = deck_title(prs)
    ctx = dict(img_dir=img_dir, rel=rel, fallback=fallback)
    chunks, report = [], []
    for i, s in enumerate(prs.slides, 1):
        if hidden(s):
            continue
        if i == 1 and title:
            report.append((i, "skipped", "deck title slide")); continue
        md, kind, why = slide_md(s, i, ctx)
        chunks.append(md); report.append((i, kind, why))
    getattr(fallback, "close", lambda: None)()
    front = (f'---\ntitle: "{(title or Path(deck).stem).replace(chr(34), chr(39))}"\n'
             'subtitle: "BASC0005 Quantitative Methods 2: Data Science and Visualisation"\n'
             'author: "Ollie Ballinger"\n---\n')
    (Path(out_dir) / f"{slug}.qmd").write_text(front + "\n" + "\n\n".join(chunks) + "\n")
    return report


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("deck"); ap.add_argument("pdf"); ap.add_argument("out_dir"); ap.add_argument("slug")
    a = ap.parse_args()
    rep = convert(a.deck, a.out_dir, a.slug, pdf=a.pdf)
    n_img = sum(1 for r in rep if r[1] == "image")
    print(f"{a.slug}: {len(rep)} slides, {len(rep) - n_img} native, {n_img} as images")
    for i, kind, why in rep:
        if kind == "image":
            print(f"   slide {i}: {why}")
